"""Turn an uploaded prescription (text, image, PDF, DOCX) into plain text."""
import io
import logging
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageOps

from . import config, llm, prompts

log = logging.getLogger(__name__)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
TEXT_EXTS = {".txt", ".md", ".csv"}
MAX_IMAGE_SIDE = 1600
MIN_PDF_PAGE_CHARS = 40  # below this a PDF page is treated as a scan and OCR'd


class ExtractionError(ValueError):
    """The file could not be read."""


@dataclass
class Extraction:
    text: str
    method: str
    pages: int = 1
    warnings: list = field(default_factory=list)


def extract(filename: str, data: bytes) -> Extraction:
    ext = Path(filename or "").suffix.lower()
    if ext in TEXT_EXTS:
        return Extraction(data.decode("utf-8", errors="replace"), "text file")
    if ext == ".pdf":
        return _from_pdf(data)
    if ext == ".docx":
        return _from_docx(data)
    if ext in IMAGE_EXTS or _looks_like_image(data):
        text, method = _read_image(_prepare_image(data))
        return Extraction(text, method, warnings=_ocr_warnings(text))
    raise ExtractionError(f"Unsupported file type '{ext or 'unknown'}'. Use an image, PDF, DOCX or TXT file.")


def _looks_like_image(data: bytes) -> bool:
    try:
        Image.open(io.BytesIO(data)).verify()
        return True
    except Exception:
        return False


def _prepare_image(data: bytes) -> bytes:
    """Fix phone-photo rotation, shrink large images and re-encode as JPEG."""
    try:
        img = Image.open(io.BytesIO(data))
        img = ImageOps.exif_transpose(img).convert("RGB")
    except Exception as exc:
        raise ExtractionError("This image could not be opened.") from exc
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=90)
    return out.getvalue()


def _read_image(jpeg: bytes) -> tuple[str, str]:
    """Read an image with the Llama vision model; fall back to Tesseract OCR."""
    if config.VISION_MODEL.lower() in {"", "none"}:
        text = _tesseract(jpeg)
        if text is None:
            raise ExtractionError(
                "Reading photos needs a vision model, and none is set up (VISION_MODEL=none). "
                "Please type the prescription text, or upload a digital PDF or Word file."
            )
        return text, "Tesseract OCR"
    try:
        messages = [{"role": "user", "content": prompts.VISION_TRANSCRIBE}]
        text = llm.complete(messages, model=config.VISION_MODEL, images=[jpeg], temperature=0)
        return text.strip(), f"vision model ({config.VISION_MODEL})"
    except llm.LLMError as exc:
        text = _tesseract(jpeg)
        if text is None:
            raise
        log.warning("Vision model unavailable (%s); used Tesseract OCR", exc)
        return text, "Tesseract OCR"


def _tesseract(jpeg: bytes):
    if not shutil.which("tesseract"):
        return None
    try:
        import pytesseract
    except ImportError:
        return None
    img = Image.open(io.BytesIO(jpeg))
    return pytesseract.image_to_string(img, lang="eng").strip()


def _ocr_warnings(text: str) -> list:
    warnings = []
    illegible = text.count("[illegible]") + text.count("(?)")
    if illegible:
        warnings.append(
            f"{illegible} word(s) could not be read with certainty. Check them against your original prescription."
        )
    if len(text.strip()) < 20:
        warnings.append("Very little text was found. Try a clearer, well-lit photo taken straight on.")
    return warnings


def _from_pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This PDF could not be opened.") from exc

    parts, methods, warnings = [], set(), []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if len(text) >= MIN_PDF_PAGE_CHARS:
            methods.add("PDF text")
        else:
            text, method = _read_image(_render_pdf_page(data, i))
            methods.add(method)
            warnings += _ocr_warnings(text)
        parts.append(f"[Page {i + 1}]\n{text}")
    return Extraction("\n\n".join(parts), " + ".join(sorted(methods)), len(reader.pages), warnings)


def _render_pdf_page(data: bytes, index: int) -> bytes:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        img = pdf[index].render(scale=2).to_pil().convert("RGB")
    finally:
        pdf.close()
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=90)
    return out.getvalue()


def _from_docx(data: bytes) -> Extraction:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This Word document could not be opened.") from exc
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Extraction("\n".join(lines), "Word document")
